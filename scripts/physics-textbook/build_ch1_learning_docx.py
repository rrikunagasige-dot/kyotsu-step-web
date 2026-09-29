#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HOLE_RE = re.compile(r'(【[A-G]\d+[a-z]?[^】]*】)')
FIG_OPEN_RE = re.compile(r'^:::figure\s+id="([^"]+)"\s+source="([^"]+)"\s+app_asset="([^"]+)"$')
CHOICES_OPEN_RE = re.compile(r'^:::choices\s+id="([^"]+)"$')


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9E2F3', size='8'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)




def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    tag = qn('w:cantSplit')
    if trPr.find(tag) is None:
        trPr.append(OxmlElement('w:cantSplit'))

def set_run_font(run, size=10.5, bold=None, color=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text_with_holes(paragraph, text: str):
    token_re = re.compile(r'(【[A-G]\d+[a-z]?[^】]*】|\*\*[^*]+\*\*)')
    parts = token_re.split(text)
    for part in parts:
        if not part:
            continue
        if HOLE_RE.fullmatch(part):
            r = paragraph.add_run(part)
            set_run_font(r, 10.5, bold=True, color=(31, 78, 121))
        elif part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, 10.5, bold=True)
        else:
            r = paragraph.add_run(part)
            set_run_font(r, 10.5)

def parse_frontmatter(lines):
    if not lines or lines[0].strip() != '---':
        return {}, lines
    meta_lines=[]
    i=1
    while i < len(lines) and lines[i].strip() != '---':
        meta_lines.append(lines[i])
        i += 1
    return {'raw': '\n'.join(meta_lines)}, lines[i+1:]


def build(source: Path, figure_dir: Path, output: Path):
    text = source.read_text(encoding='utf-8')
    lines = text.splitlines()
    _, lines = parse_frontmatter(lines)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in [('Title', 22), ('Heading 1', 16), ('Heading 2', 13)]:
        st = styles[name]
        st.font.name = 'Aptos'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
        st.font.size = Pt(size)
        st.font.bold = True

    blocks=[]
    i=0
    while i < len(lines):
        line=lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith(':::figure'):
            m=FIG_OPEN_RE.match(line)
            if not m:
                raise ValueError(f'Bad figure directive: {line}')
            fig_id, src_name, app_asset = m.groups()
            i += 1
            content=[]
            while i < len(lines) and lines[i].strip() != ':::':
                content.append(lines[i].rstrip())
                i += 1
            blocks.append(('figure', fig_id, src_name, app_asset, ' '.join(content).strip()))
            i += 1
            continue
        if line.startswith(':::choices'):
            m=CHOICES_OPEN_RE.match(line)
            if not m:
                raise ValueError(f'Bad choices directive: {line}')
            qid=m.group(1)
            i += 1
            content=[]
            while i < len(lines) and lines[i].strip() != ':::':
                content.append(lines[i].rstrip())
                i += 1
            blocks.append(('choices', qid, ' '.join(content).strip()))
            i += 1
            continue
        if line.strip() == ':::callout':
            i += 1
            content=[]
            while i < len(lines) and lines[i].strip() != ':::':
                content.append(lines[i].rstrip())
                i += 1
            blocks.append(('callout', ' '.join(content).strip()))
            i += 1
            continue
        if line.startswith('# '):
            blocks.append(('h1', line[2:].strip()))
        elif line.startswith('## '):
            blocks.append(('h2', line[3:].strip()))
        else:
            blocks.append(('p', line.strip()))
        i += 1

    # First h1 gets title treatment if it is the cover chapter title; preserve later h1s.
    h1_count=0
    fig_count=0
    for idx, block in enumerate(blocks):
        kind=block[0]
        if kind=='h1':
            h1_count += 1
            txt=block[1]
            if h1_count==1:
                p=doc.add_paragraph(style='Title')
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(txt); set_run_font(r,22,bold=True)
            else:
                p=doc.add_paragraph(style='Heading 1')
                r=p.add_run(txt); set_run_font(r,16,bold=True)
        elif kind=='h2':
            p=doc.add_paragraph(style='Heading 2')
            r=p.add_run(block[1]); set_run_font(r,13,bold=True)
        elif kind=='p':
            p=doc.add_paragraph()
            add_text_with_holes(p, block[1])
            if idx + 1 < len(blocks) and blocks[idx+1][0] == 'choices':
                p.paragraph_format.keep_with_next = True
        elif kind=='callout':
            table=doc.add_table(rows=1, cols=1)
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            cell=table.cell(0,0)
            set_cell_shading(cell, 'F3F6FA')
            set_cell_border(cell, color='AFC4DD')
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]
            add_text_with_holes(p, block[1])
        elif kind=='choices':
            _, qid, content=block
            table=doc.add_table(rows=1, cols=1)
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            prevent_row_split(table.rows[0])
            cell=table.cell(0,0)
            set_cell_shading(cell, 'F8FBFF')
            set_cell_border(cell, color='C5D9F1')
            p=cell.paragraphs[0]
            r=p.add_run(f'選択肢  {qid}\n'); set_run_font(r,10,bold=True,color=(31,78,121))
            r=p.add_run(content); set_run_font(r,10)
        elif kind=='figure':
            _, fig_id, src_name, app_asset, caption=block
            fig_path=figure_dir/src_name
            if not fig_path.exists():
                raise FileNotFoundError(fig_path)
            p=doc.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run()
            run.add_picture(str(fig_path), width=Cm(13.5))
            cp=doc.add_paragraph()
            cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            rr=cp.add_run(caption); set_run_font(rr,9)
            fig_count += 1

    # Footer marker
    for section in doc.sections:
        p=section.footer.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('Chapter 1 learning-text checkpoint — 2026-09-29')
        set_run_font(r,8,color=(100,100,100))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    hole_ids=re.findall(r'【([A-G]\d+[a-z]?)', text)
    print(f'output={output}')
    print(f'hole_occurrences={len(hole_ids)} unique_holes={len(set(hole_ids))} figures={fig_count}')


def main():
    ap=argparse.ArgumentParser(description='Build Chapter-1 physics learning-text review DOCX from the authoritative markdown source.')
    ap.add_argument('--source', type=Path, required=True)
    ap.add_argument('--figure-dir', type=Path, required=True)
    ap.add_argument('--output', type=Path, required=True)
    args=ap.parse_args()
    build(args.source, args.figure_dir, args.output)


if __name__=='__main__':
    main()